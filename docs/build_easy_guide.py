from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).with_name("Guide-simple-Malagasy-Events.docx")

RED = "C8102E"
GREEN = "007A3D"
PURPLE = "3C3489"
GOLD = "B8860B"
BLUE = "1565C0"
INK = "202124"
GRAY = "666666"
WHITE = "FFFFFF"
LIGHT = "F6F7F8"
PALE_GREEN = "EAF6EF"
PALE_RED = "FDECEF"
PALE_PURPLE = "F0EDFF"
PALE_GOLD = "FFF5D6"
PALE_BLUE = "EAF2FD"


def set_font(run, size=12, color=INK, bold=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def shade_paragraph(paragraph, fill, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "28")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("MALAGASY EVENTS  •  ")
    set_font(r, 8, GREEN, True)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, text, end])


def add_title(doc, title, kicker=None):
    if kicker:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(kicker.upper())
        set_font(r, 9, GREEN, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    set_font(r, 26, INK, True)
    return p


def add_intro(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    set_font(r, 13, GRAY)
    return p


def add_heading(doc, text, color=GREEN):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, 17, color, True)
    return p


def add_bullet(doc, text, color=INK):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_font(r, 12, color)
    return p


def new_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_nums = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abs, default=0) + 1
    num_id = max(existing_nums, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.extend([tabs, ind])
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_step(doc, text, num_id):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_font(r, 12, INK)
    return p


def add_box(doc, title, text, fill, accent):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.10)
    p.paragraph_format.right_indent = Inches(0.10)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.22
    shade_paragraph(p, fill, accent)
    r = p.add_run(title + "\n")
    set_font(r, 13, accent, True)
    r2 = p.add_run(text)
    set_font(r2, 11.5, INK)
    return p


def add_yes_no(doc, yes, no):
    add_box(doc, "✓ Tu peux", yes, PALE_GREEN, GREEN)
    add_box(doc, "× Tu ne peux pas encore", no, PALE_RED, RED)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.85)
sec.bottom_margin = Inches(0.8)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.2

for style_name in ("List Bullet", "List Number"):
    st = doc.styles[style_name]
    st.font.name = "Calibri"
    st.font.size = Pt(12)
    st.paragraph_format.left_indent = Inches(0.42)
    st.paragraph_format.first_line_indent = Inches(-0.22)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.2

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("GUIDE FACILE  |  MALAGASY EVENTS")
set_font(hr, 8, GREEN, True)
add_page_number(sec.footer.paragraphs[0])

# Couverture
doc.add_paragraph().paragraph_format.space_after = Pt(58)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MALAGASY")
set_font(r, 20, RED, True)
r = p.add_run(" EVENTS")
set_font(r, 20, GREEN, True)

doc.add_paragraph().paragraph_format.space_after = Pt(30)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run("Le guide\nsuper facile")
set_font(r, 34, INK, True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(30)
r = p.add_run("Je regarde. Je clique. Je participe.")
set_font(r, 15, GRAY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(54)
r = p.add_run("PUBLIC  •  INSCRIT  •  PREMIUM  •  ORGANISATEUR")
set_font(r, 11, GREEN, True)

add_box(
    doc,
    "Le but de ce guide",
    "T’aider à comprendre le site sans mots compliqués. Tu peux suivre les étapes dans l’ordre, comme une recette.",
    LIGHT,
    GREEN,
)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Version simple • Juillet 2026")
set_font(r, 9, GRAY)

# Page 2
doc.add_page_break()
add_title(doc, "D’abord, qui es-tu ?", "Étape 1")
add_intro(doc, "Choisis la phrase qui te ressemble. Ensuite, va directement à la bonne page du guide.")
add_box(doc, "Je visite seulement", "Tu es PUBLIC. Tu n’as pas besoin de compte pour regarder les événements et les bonnes adresses.", PALE_GREEN, GREEN)
add_box(doc, "J’ai créé un compte gratuit", "Tu es INSCRIT. Tu peux participer, suivre des personnes et recevoir des nouvelles.", PALE_RED, RED)
add_box(doc, "J’ai le forfait Premium", "Tu es PREMIUM. Tu as les fonctions du compte gratuit, avec des avantages en plus.", PALE_PURPLE, PURPLE)
add_box(doc, "J’organise des événements", "Tu es ORGANISATEUR. Tu peux gérer ta fiche et publier pour ton association ou ton entreprise.", PALE_GOLD, GOLD)
add_heading(doc, "Une règle très simple")
add_box(doc, "Plus tu as de droits, plus tu peux faire de choses", "Premium garde tout ce qu’un inscrit peut faire. Organisateur garde aussi les fonctions d’un membre.", PALE_BLUE, BLUE)

# Page 3
doc.add_page_break()
add_title(doc, "Je suis PUBLIC", "Sans compte")
add_intro(doc, "Tu peux déjà découvrir presque tout le site.")
add_yes_no(
    doc,
    "Voir les événements, chercher une ville, utiliser les filtres, ouvrir le calendrier, consulter les restaurants, boutiques, églises, sportifs, professionnels, la Diaspora et le Guide France.",
    "Écrire dans la Communauté, envoyer des messages, commenter, suivre une personne ou voir les After-movies."
)
add_heading(doc, "Je cherche un événement")
steps_public = new_numbering(doc)
add_step(doc, "Ouvre la page Événements.", steps_public)
add_step(doc, "Écris un mot : Paris, Lyon, concert, sport…", steps_public)
add_step(doc, "Choisis l’événement qui te plaît.", steps_public)
add_step(doc, "Lis la date, le lieu et le prix.", steps_public)
add_step(doc, "S’il existe une vraie billetterie vérifiée, clique sur « Acheter mes billets ».", steps_public)
add_box(doc, "Pas de vraie billetterie ?", "Le bouton reste gris et dit que la billetterie arrive bientôt. Tu ne seras jamais envoyé vers un faux lien.", PALE_GOLD, GOLD)

# Page 4
doc.add_page_break()
add_title(doc, "J’ai un compte GRATUIT", "Inscrit")
add_intro(doc, "Tu peux faire tout ce que fait le public, puis participer à la vie du site.")
add_heading(doc, "Avec mon compte, je peux…")
for text in [
    "Dire que je suis intéressé par un événement.",
    "Ajouter un événement à mon calendrier.",
    "Demander un rappel.",
    "Écrire un commentaire.",
    "Publier un message ou une photo dans la Communauté.",
    "Aimer et commenter les publications.",
    "Suivre des membres et des organisateurs.",
    "Envoyer des messages privés.",
    "Voir les After-movies.",
    "Proposer un événement à l’équipe.",
]:
    add_bullet(doc, text)
add_heading(doc, "Je crée mon compte")
steps_account = new_numbering(doc)
add_step(doc, "Clique sur « Connexion ».", steps_account)
add_step(doc, "Choisis « S’inscrire ».", steps_account)
add_step(doc, "Écris ton e-mail et ton mot de passe.", steps_account)
add_step(doc, "Choisis un pseudo qui n’est pas déjà utilisé.", steps_account)
add_step(doc, "Ajoute une photo si tu veux.", steps_account)
add_box(doc, "Important", "Ton pseudo est unique. Deux personnes ne peuvent pas avoir exactement le même.", PALE_BLUE, BLUE)

# Page 5
doc.add_page_break()
add_title(doc, "Je suis PREMIUM", "Forfait membre")
add_intro(doc, "Premium sert à soutenir le projet et à recevoir quelques avantages en plus.")
add_heading(doc, "Je garde toutes les fonctions gratuites")
add_bullet(doc, "Événements, Communauté, messages, commentaires, suivi et notifications.")
add_heading(doc, "Et je reçois aussi…", PURPLE)
for text in [
    "Un badge Premium visible sur mon profil.",
    "Mes publications mises en avant pendant 48 heures.",
    "Des réductions chez les partenaires actifs.",
    "Un QR Premium sécurisé pour utiliser certaines offres.",
    "Une priorité possible sur les billets et tombolas partenaires.",
]:
    add_bullet(doc, text)
add_box(doc, "Prix affiché", "2,50 € par mois, sans engagement.", PALE_PURPLE, PURPLE)
add_box(doc, "Attention", "Un avantage apparaît seulement quand un vrai partenaire l’a activé. Premium ne transforme pas ton compte en compte Organisateur.", PALE_GOLD, GOLD)

# Page 6
doc.add_page_break()
add_title(doc, "Je suis ORGANISATEUR", "Association ou professionnel")
add_intro(doc, "Ce compte sert à parler au nom d’une structure : association, club, artiste, organisateur ou entreprise.")
add_heading(doc, "Je peux…", GOLD)
for text in [
    "Créer une fiche pour ma structure.",
    "Demander le contrôle d’une fiche qui existe déjà.",
    "Ajouter mon logo, ma ville et mes liens.",
    "Publier mes événements directement.",
    "Publier une actualité au nom de ma structure.",
    "Voir mes abonnés et mes statistiques de base.",
    "Répondre aux messages reçus.",
    "Ajouter des personnes à mon équipe.",
]:
    add_bullet(doc, text)
add_heading(doc, "Je publie un événement")
steps_publish = new_numbering(doc)
add_step(doc, "Je me connecte avec mon compte Organisateur.", steps_publish)
add_step(doc, "Je clique sur « Proposer un événement ».", steps_publish)
add_step(doc, "J’écris le nom, la date, le lieu, le prix et la description.", steps_publish)
add_step(doc, "J’ajoute une affiche et une vraie billetterie si elles existent.", steps_publish)
add_step(doc, "Je vérifie toutes les informations.", steps_publish)
add_step(doc, "Je publie.", steps_publish)
add_box(doc, "Règle importante", "Un événement doit être rattaché au bon organisateur. On ne met jamais l’événement d’une autre structure sur sa fiche.", PALE_RED, RED)

# Page 7
doc.add_page_break()
add_title(doc, "À quoi servent les onglets ?", "La carte du site")
add_intro(doc, "Chaque onglet est comme une pièce différente dans une grande maison.")
items = [
    ("Événements", "Les sorties, concerts, soirées, fêtes et tournois."),
    ("Diaspora", "Une présentation de la communauté malagasy en France."),
    ("Gastronomie", "Restaurants, traiteurs et food trucks."),
    ("Professionnels", "Associations, organisateurs, artistes, médias et groupes."),
    ("Églises", "Paroisses et communautés chrétiennes malagasy."),
    ("Sportifs", "Clubs, associations et événements sportifs."),
    ("Boutiques", "Épiceries, vêtements, créations et artisanat."),
    ("Guide France", "Des aides pour arriver, étudier et travailler en France."),
    ("Communauté", "Les publications, échanges et messages entre membres."),
    ("After-movies", "Les vidéos et souvenirs des événements."),
    ("Mes abonnements", "La page qui explique Premium et Organisateur."),
    ("Notifications", "Les nouvelles activités qui te concernent."),
]
for title, desc in items:
    add_box(doc, title, desc, LIGHT, GREEN if title not in ("Communauté", "After-movies") else RED)

# Page 8
doc.add_page_break()
add_title(doc, "Mes actions les plus utiles", "Petites recettes")
add_heading(doc, "Je veux suivre un organisateur")
steps_follow = new_numbering(doc)
add_step(doc, "Je crée un compte ou je me connecte.", steps_follow)
add_step(doc, "J’ouvre l’onglet Professionnels.", steps_follow)
add_step(doc, "Je cherche son nom.", steps_follow)
add_step(doc, "J’ouvre sa fiche.", steps_follow)
add_step(doc, "Je clique sur « Suivre ».", steps_follow)
add_heading(doc, "Je veux parler avec la communauté")
steps_community = new_numbering(doc)
add_step(doc, "Je me connecte.", steps_community)
add_step(doc, "J’ouvre l’onglet Communauté.", steps_community)
add_step(doc, "J’écris mon message ou j’ajoute une photo.", steps_community)
add_step(doc, "Je vérifie que mon message est gentil et utile.", steps_community)
add_step(doc, "Je clique sur « Publier ».", steps_community)
add_heading(doc, "Je veux trouver une bonne adresse")
steps_address = new_numbering(doc)
add_step(doc, "J’ouvre Gastronomie, Boutiques ou Églises.", steps_address)
add_step(doc, "Je choisis un type ou une région.", steps_address)
add_step(doc, "J’ouvre la fiche.", steps_address)
add_step(doc, "Je regarde la carte et les contacts.", steps_address)

# Page 9
doc.add_page_break()
add_title(doc, "Je reste en sécurité", "Très important")
add_intro(doc, "Le site est fait pour se rencontrer et s’entraider. Mais chacun doit rester prudent.")
for text in [
    "Je ne donne jamais mon mot de passe.",
    "Je vérifie le nom, la date et le lieu avant de payer.",
    "Je clique seulement sur une billetterie vérifiée.",
    "Je reste poli dans les messages et les commentaires.",
    "Je peux bloquer une personne qui me dérange.",
    "Je peux signaler un contenu dangereux ou méchant.",
]:
    add_bullet(doc, text)
add_box(doc, "Si quelque chose semble bizarre", "Ne paie pas. Ne réponds pas. Fais une capture d’écran et contacte l’équipe Malagasy Events.", PALE_RED, RED)
add_heading(doc, "Ce qui arrive plus tard")
for text in [
    "La billetterie complète de Malagasy Events.",
    "Davantage d’avantages Premium.",
    "Des statistiques plus détaillées pour les organisateurs.",
    "Plus de partenaires pour le Guide France.",
]:
    add_bullet(doc, text, GRAY)
add_box(doc, "Tu as compris l’essentiel !", "Regarde, cherche, clique et participe. Si une fonction demande un compte, le site te proposera simplement de te connecter.", PALE_GREEN, GREEN)

doc.core_properties.title = "Le guide super facile — Malagasy Events"
doc.core_properties.subject = "Guide simple pour comprendre les fonctions du site"
doc.core_properties.author = "Malagasy Events"
doc.core_properties.keywords = "Malagasy Events, guide facile, public, inscrit, premium, organisateur"
doc.save(OUT)
print(OUT)
