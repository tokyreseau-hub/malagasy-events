from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


OUT = Path(__file__).with_name("Guide-fonctionnalites-Malagasy-Events.docx")

RED = "C8102E"
GREEN = "007A3D"
PURPLE = "3C3489"
GOLD = "B8860B"
INK = "202124"
MUTED = "666666"
LIGHT = "F5F6F7"
PALE_GREEN = "EAF6EF"
PALE_RED = "FDECEF"
PALE_PURPLE = "F0EDFF"
PALE_GOLD = "FFF7DA"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, *, bold=False, color=INK, size=8.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)


def set_no_cell_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("MALAGASY EVENTS  •  ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(GREEN)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_title(doc, title, kicker=None):
    if kicker:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(kicker.upper())
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(GREEN)
    p = doc.add_paragraph(style="Title")
    p.add_run(title)


def add_intro(doc, text):
    p = doc.add_paragraph(text)
    p.style = doc.styles["Lead"]
    p.paragraph_format.space_after = Pt(12)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_bullet(doc, text, *, color=INK):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_callout(doc, label, text, fill=PALE_GREEN, accent=GREEN):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    set_no_cell_borders(table)
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label.upper())
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    for r2 in p2.runs:
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_role_cards(doc):
    table = doc.add_table(rows=1, cols=4)
    widths = [2340] * 4
    set_table_width(table, widths)
    set_no_cell_borders(table)
    roles = [
        ("PUBLIC", "Sans compte", "Consulter, chercher, filtrer et partager les contenus publics.", PALE_GREEN, GREEN),
        ("INSCRIT", "Compte gratuit", "Interagir, publier, suivre, commenter, recevoir des notifications.", PALE_RED, RED),
        ("PREMIUM", "2,50 € / mois", "Profiter des avantages et d’une visibilité renforcée.", PALE_PURPLE, PURPLE),
        ("ORGANISATEUR", "15 € / mois", "Gérer une fiche, publier directement et animer sa communauté.", PALE_GOLD, GOLD),
    ]
    for i, (name, sub, desc, fill, accent) in enumerate(roles):
        c = table.cell(0, i)
        shade(c, fill)
        set_cell_margins(c, top=150, start=140, bottom=150, end=140)
        set_cell_text(c, name, bold=True, color=accent, size=9.5)
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(sub)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(INK)
        p2 = c.add_paragraph(desc)
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.08
        for r2 in p2.runs:
            r2.font.size = Pt(8.3)
            r2.font.color.rgb = RGBColor.from_string(MUTED)


def add_access_matrix(doc):
    rows = [
        ("Événements", "Lire / filtrer", "Interagir / proposer", "Même accès + badge", "Publier / gérer"),
        ("Diaspora", "Accès complet", "Accès complet", "Accès complet", "Accès complet"),
        ("Gastronomie", "Annuaire / carte", "Annuaire / carte", "Avantages partenaires", "Référencement pro"),
        ("Professionnels", "Annuaire / fiches", "Suivre / contacter", "Accès complet", "Gérer sa fiche"),
        ("Églises", "Annuaire / carte", "Accès complet", "Accès complet", "Référencement"),
        ("Sportifs", "Clubs + événements", "Suivre / interagir", "Accès complet", "Gérer sa fiche"),
        ("Boutiques", "Annuaire / carte", "Accès complet", "Réductions partenaires", "Fiche Premium possible"),
        ("Guide France", "Accès complet", "Accès complet", "Accès complet", "Accès complet"),
        ("After-movies", "Connexion requise", "Voir les vidéos", "Voir les vidéos", "Voir / contribuer"),
        ("Communauté", "Connexion requise", "Publier / échanger", "Mise en avant 48 h", "Publier comme orga"),
        ("Profil & alertes", "Connexion requise", "Profil / messages", "Badge Premium", "Équipe / statistiques"),
        ("Offres", "Voir les offres", "Souscrire", "Gérer son statut", "Accéder à l’offre"),
    ]
    table = doc.add_table(rows=1, cols=5)
    widths = [1900, 1765, 1895, 1850, 1950]
    set_table_width(table, widths)
    table.style = "Table Grid"
    headers = ["Onglet", "Public", "Inscrit", "Premium", "Organisateur"]
    colors = [INK, GREEN, RED, PURPLE, GOLD]
    for i, h in enumerate(headers):
        shade(table.cell(0, i), colors[i])
        set_cell_text(table.cell(0, i), h, bold=True, color=WHITE, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            shade(cells[i], WHITE if idx % 2 == 0 else LIGHT)
            set_cell_text(cells[i], text, bold=(i == 0), size=8.1)


def add_access_block(doc, title, public, registered, premium, orga, note=None):
    add_heading(doc, title, 2)
    table = doc.add_table(rows=4, cols=2)
    set_table_width(table, [1550, 7810])
    set_no_cell_borders(table)
    items = [
        ("PUBLIC", public, PALE_GREEN, GREEN),
        ("INSCRIT", registered, PALE_RED, RED),
        ("PREMIUM", premium, PALE_PURPLE, PURPLE),
        ("ORGANISATEUR", orga, PALE_GOLD, GOLD),
    ]
    for r, (label, text, fill, accent) in enumerate(items):
        shade(table.cell(r, 0), accent)
        shade(table.cell(r, 1), fill)
        set_cell_text(table.cell(r, 0), label, bold=True, color=WHITE, size=8.5)
        set_cell_text(table.cell(r, 1), text, size=9.2)
    if note:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run("À noter — ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(RED)
        r2 = p.add_run(note)
        r2.font.color.rgb = RGBColor.from_string(MUTED)
        r2.font.size = Pt(9)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.8)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.36)
sec.footer_distance = Inches(0.38)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(5)

styles.add_style("Lead", WD_STYLE_TYPE.PARAGRAPH)
styles["Lead"].font.name = "Calibri"
styles["Lead"].font.size = Pt(12)
styles["Lead"].font.color.rgb = RGBColor.from_string(MUTED)
styles["Lead"].paragraph_format.line_spacing = 1.2

styles["Title"].font.name = "Calibri"
styles["Title"].font.size = Pt(25)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor.from_string(INK)
styles["Title"].paragraph_format.space_after = Pt(8)

for name, size, color in (("Heading 1", 17, RED), ("Heading 2", 12.5, GREEN), ("Heading 3", 10.5, PURPLE)):
    st = styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

styles["List Bullet"].font.name = "Calibri"
styles["List Bullet"].font.size = Pt(10)
styles["List Bullet"].paragraph_format.left_indent = Inches(0.22)
styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.16)

# Header / footer
header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("MALAGASY EVENTS  |  GUIDE FONCTIONNEL")
hr.bold = True
hr.font.name = "Calibri"
hr.font.size = Pt(8)
hr.font.color.rgb = RGBColor.from_string(GREEN)
add_page_number(sec.footer.paragraphs[0])

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(36)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MALAGASY")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = RGBColor.from_string(RED)
r2 = p.add_run(" EVENTS")
r2.bold = True
r2.font.size = Pt(20)
r2.font.color.rgb = RGBColor.from_string(GREEN)

doc.add_paragraph().paragraph_format.space_after = Pt(46)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
r = p.add_run("Guide des fonctionnalités\net des niveaux d’accès")
r.bold = True
r.font.name = "Calibri"
r.font.size = Pt(31)
r.font.color.rgb = RGBColor.from_string(INK)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(28)
r = p.add_run("Public  •  Inscrit  •  Premium  •  Organisateur")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor.from_string(MUTED)

cover_table = doc.add_table(rows=1, cols=3)
set_table_width(cover_table, [3120, 3120, 3120])
set_no_cell_borders(cover_table)
for i, color in enumerate((WHITE, RED, GREEN)):
    shade(cover_table.cell(0, i), color)
    cover_table.cell(0, i).height = Inches(0.12)

doc.add_paragraph().paragraph_format.space_after = Pt(60)
add_callout(
    doc,
    "Objet du document",
    "Présenter, onglet par onglet, ce que chaque type d’utilisateur peut consulter, faire et gérer sur Malagasy Events. Les fonctions encore en préparation sont signalées pour éviter toute confusion.",
    fill=LIGHT,
    accent=GREEN,
)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(22)
r = p.add_run("Version 1.0  •  26 juillet 2026  •  France")
r.font.size = Pt(9)
r.font.color.rgb = RGBColor.from_string(MUTED)

# Page 2
doc.add_page_break()
add_title(doc, "Les quatre niveaux d’accès", "Mode d’emploi")
add_intro(doc, "Les droits s’additionnent : un compte Premium conserve toutes les fonctions d’un compte inscrit ; un compte Organisateur dispose des fonctions membres et ajoute les outils de gestion professionnelle.")
add_role_cards(doc)
doc.add_paragraph()
add_heading(doc, "Repères utilisés dans ce guide", 1)
add_bullet(doc, "Disponible : la fonction est utilisable dans la version actuelle du site.")
add_bullet(doc, "Connexion requise : le contenu ou l’action est réservé aux membres.")
add_bullet(doc, "Réservé Premium / Organisateur : l’accès dépend du forfait actif.")
add_bullet(doc, "Bientôt : l’interface ou l’avantage est annoncé, mais dépend encore d’une activation, d’un paiement ou d’un partenaire.")
add_callout(doc, "Principe général", "Les pages d’information et les annuaires restent largement publics pour favoriser le référencement Google. Les fonctions sociales, personnelles et de gestion demandent une connexion.", PALE_GREEN, GREEN)

# Page 3
doc.add_page_break()
add_title(doc, "Vue d’ensemble par onglet", "Matrice des accès")
add_intro(doc, "Cette matrice donne le niveau d’action principal. Les pages suivantes détaillent les fonctions et les limites de chaque onglet.")
add_access_matrix(doc)

# Page 4
doc.add_page_break()
add_title(doc, "Événements & agenda", "Onglet principal")
add_intro(doc, "Le cœur de Malagasy Events : découvrir les rendez-vous de la diaspora, trouver rapidement une date et passer de la découverte à l’action.")
add_access_block(
    doc, "Accueil / Événements",
    "Voir les événements à venir ou passés, rechercher par mot-clé, filtrer par ville et catégorie, afficher la liste ou le calendrier, ouvrir une fiche événement et partager son lien.",
    "Toutes les fonctions publiques, puis indiquer « Intéressé », commenter, demander un rappel, ajouter à Google Agenda et proposer un événement soumis à validation.",
    "Toutes les fonctions du membre ; le badge Premium est visible dans les interactions. Les publications communautaires peuvent être remontées pendant 48 heures.",
    "Publier directement ses événements depuis son identité organisateur, relier les événements à sa fiche, suivre l’intérêt et les commentaires, et bénéficier d’une visibilité renforcée selon l’offre.",
    "La billetterie interne et le paiement en ligne ne sont pas encore actifs. Un lien externe de billetterie peut être affiché lorsqu’il existe."
)
add_heading(doc, "Actions disponibles sur une fiche événement", 1)
for t in [
    "Consulter la date, la ville, le lieu, le tarif, l’affiche, la description et l’organisateur.",
    "Ouvrir les liens externes utiles : billetterie, réseaux ou informations complémentaires.",
    "Partager l’événement et l’ajouter à Google Agenda.",
    "Interagir après connexion : intérêt, commentaire, rappel et entraide liée à l’événement.",
]:
    add_bullet(doc, t)
add_callout(doc, "Visibilité", "Les événements mis « À la une » sont visibles de tous. Les compteurs et interactions sociales sont associés aux comptes connectés.", PALE_RED, RED)

# Page 5
doc.add_page_break()
add_title(doc, "Annuaires de la communauté", "Trouver une structure ou une adresse")
add_intro(doc, "Les annuaires sont publics afin d’aider la communauté et d’améliorer la découverte sur Google. Les membres ajoutent ensuite le suivi et la prise de contact.")
add_access_block(
    doc, "Gastronomie",
    "Explorer restaurants, traiteurs, food trucks, guides et communautés ; filtrer par type ou région ; voir les adresses sur la carte et ouvrir les réseaux ou contacts.",
    "Même accès, avec la possibilité de poursuivre les échanges depuis les espaces membres.",
    "Afficher les avantages partenaires disponibles et présenter un QR Premium sécurisé chez le partenaire.",
    "Référencer ou enrichir une activité ; une fiche annuaire Premium peut être épinglée et valorisée selon l’offre.",
    "Les réductions n’apparaissent que lorsqu’un partenariat actif a été enregistré."
)
add_access_block(
    doc, "Professionnels & associations",
    "Parcourir les associations, organisateurs, artistes, médias, groupes et autres structures ; filtrer et ouvrir chaque fiche.",
    "Suivre une structure, consulter ses événements et actualités, puis la contacter lorsqu’un moyen de contact est disponible.",
    "Même accès, avec les avantages Premium applicables au compte membre.",
    "Réclamer une fiche existante ou créer sa structure, personnaliser sa présentation et publier au nom de l’organisme."
)
add_access_block(
    doc, "Églises",
    "Consulter les paroisses et communautés chrétiennes malagasy, filtrer, utiliser la carte et ouvrir les contacts disponibles.",
    "Même accès public, enrichi des fonctions générales du compte.",
    "Même accès que le membre.",
    "Être référencé comme structure ; la gestion directe dépend du rattachement de la fiche."
)

# Page 6
doc.add_page_break()
add_title(doc, "Sport, boutiques & artisanat", "Annuaires spécialisés")
add_access_block(
    doc, "Sportifs",
    "Découvrir les clubs, associations sportives et organisateurs de tournois ; voir uniquement les événements sportifs associés.",
    "Suivre les structures, ouvrir leurs fiches et interagir avec les événements sportifs.",
    "Même accès que le membre, avec le statut Premium visible dans les interactions.",
    "Gérer sa fiche sportive et rattacher uniquement ses propres événements.",
    "L’association entre une structure et un événement repose sur le bon rattachement de l’organisateur."
)
add_access_block(
    doc, "Boutiques & artisanat",
    "Découvrir les boutiques, épiceries, créateurs et artisans malagasy ; filtrer, voir la carte et utiliser les coordonnées disponibles.",
    "Même accès, avec les fonctions de compte et de contact.",
    "Profiter des réductions partenaires actives lorsqu’elles existent.",
    "Référencer son activité ; souscrire éventuellement une Fiche Premium annuelle pour être épinglé et mieux présenté."
)
add_heading(doc, "Fiche Premium d’annuaire", 1)
for t in [
    "Tarif affiché : 79 € / an.",
    "Fiche épinglée en tête d’annuaire.",
    "Badge doré, mise en avant sur l’accueil, photo et description enrichies.",
    "Statistiques de visite annoncées dans l’offre.",
]:
    add_bullet(doc, t)

# Page 7
doc.add_page_break()
add_title(doc, "Diaspora & Guide France", "Informations utiles")
add_access_block(
    doc, "Diaspora",
    "Lire la présentation de la diaspora malagasy en France, comprendre les termes Malagasy, malgache et gasy, et accéder aux principales rubriques du site.",
    "Même accès que le public.",
    "Même accès que le public.",
    "Même accès que le public ; cette page sert de porte d’entrée éditoriale et SEO."
)
add_access_block(
    doc, "Guide France",
    "Accéder aux quatre parcours : étudiant arrivant en France, recherche d’emploi en fin d’études, arrivée professionnelle et aide pour les personnes sans papiers.",
    "Même accès que le public ; les ressources peuvent être consultées sans compte.",
    "Même accès que le public.",
    "Même accès ; de futurs partenaires pourront enrichir les offres d’emploi et l’accompagnement professionnel.",
    "Certaines sections sont volontairement présentées comme des bases à compléter avec des partenaires spécialisés."
)
add_heading(doc, "Pages d’information", 1)
table = doc.add_table(rows=4, cols=2)
set_table_width(table, [2400, 6960])
set_no_cell_borders(table)
for i, (a, b) in enumerate([
    ("À propos", "Mission, positionnement et fonctionnement de Malagasy Events."),
    ("FAQ", "Réponses sur la publication, les comptes, les annuaires et les offres."),
    ("Contact", "Contacter l’équipe et proposer un événement ; une connexion peut être demandée pour certaines actions."),
    ("Offres", "Comparer Membre Premium et Organisateur, puis ouvrir les pages de détail."),
]):
    shade(table.cell(i, 0), GREEN if i % 2 == 0 else RED)
    shade(table.cell(i, 1), LIGHT)
    set_cell_text(table.cell(i, 0), a, bold=True, color=WHITE, size=9)
    set_cell_text(table.cell(i, 1), b, size=9.2)

# Page 8
doc.add_page_break()
add_title(doc, "Communauté & After-movies", "Espaces réservés aux membres")
add_access_block(
    doc, "Communauté",
    "Voir l’existence de l’onglet, mais se connecter pour accéder au contenu et aux échanges.",
    "Publier du texte ou une photo, aimer, commenter et répondre ; découvrir des profils, suivre des membres, envoyer des messages privés et utiliser l’entraide liée aux événements.",
    "Toutes les fonctions membres, avec badge Premium et publications remontées en tête du fil pendant 48 heures.",
    "Publier soit avec son identité personnelle, soit au nom de son organisme ; échanger avec les membres et les abonnés de la structure."
)
add_access_block(
    doc, "After-movies",
    "Onglet masqué tant que l’utilisateur n’est pas connecté.",
    "Voir les after-movies, teasers et contenus vidéo de la communauté.",
    "Même accès que le membre.",
    "Voir les contenus et contribuer selon les règles de publication et de modération."
)
add_heading(doc, "Sécurité et confort communautaire", 1)
for t in [
    "Un utilisateur peut bloquer un autre compte afin de couper l’accès direct et les échanges indésirables.",
    "Les messages, commentaires et publications sont rattachés à un profil identifiable.",
    "Les notifications informent des nouvelles interactions ; elles peuvent être marquées comme lues.",
    "L’administrateur peut être averti d’une nouvelle inscription et un message de bienvenue personnalisé peut être envoyé.",
]:
    add_bullet(doc, t)

# Page 9
doc.add_page_break()
add_title(doc, "Compte, profil & notifications", "Fonctions après connexion")
add_access_block(
    doc, "Profil personnel",
    "Connexion ou inscription depuis le bouton « Connexion ».",
    "Choisir un nom d’utilisateur unique, ajouter une photo, renseigner son profil et ses préférences, consulter abonnés et abonnements, puis se déconnecter.",
    "Afficher le badge doré Premium sur le profil, les publications, commentaires et discussions.",
    "Basculer entre identité personnelle et identité organisateur lorsqu’une fiche est rattachée."
)
add_access_block(
    doc, "Messages & notifications",
    "Connexion requise.",
    "Recevoir des notifications, suivre les messages non lus, ouvrir la messagerie privée, suivre des membres ou organisateurs et retrouver ses abonnements.",
    "Même accès que le membre.",
    "Recevoir et envoyer des messages au nom de la structure, suivre les interactions et accéder aux indicateurs de sa fiche."
)
add_heading(doc, "Règles de compte", 1)
for t in [
    "Le nom d’utilisateur est unique : deux comptes ne peuvent pas conserver le même pseudo.",
    "Les contenus personnels et sociaux sont réservés aux personnes connectées.",
    "Le menu « Mes abonnements » ouvre la page qui explique les offres Premium et Organisateur.",
    "Le français reste la langue par défaut ; les éléments fixes du site peuvent être affichés en français, anglais ou malagasy.",
]:
    add_bullet(doc, t)

# Page 10
doc.add_page_break()
add_title(doc, "Membre Premium", "Offre individuelle")
add_intro(doc, "Le Premium complète le compte gratuit pour les membres qui souhaitent soutenir la plateforme et profiter d’avantages communautaires.")
add_heading(doc, "Avantages affichés", 1)
for t in [
    "Badge doré visible sur le profil, les publications, commentaires et discussions.",
    "Publications communautaires mises en avant pendant 48 heures.",
    "Réductions exclusives chez les restaurants, boutiques et artisans partenaires.",
    "QR Premium sécurisé à présenter chez les partenaires actifs.",
    "Priorité annoncée sur les billets d’événements partenaires et les tombolas réservées.",
]:
    add_bullet(doc, t)
add_callout(doc, "Tarif affiché", "2,50 € par mois, sans engagement.", PALE_PURPLE, PURPLE)
add_callout(doc, "État actuel", "Le paiement en ligne est annoncé comme « très bientôt ». L’activation peut donc encore nécessiter un contact avec l’équipe. Les réductions, billets prioritaires et tombolas dépendent de partenaires actifs.", PALE_RED, RED)
add_heading(doc, "Ce que Premium ne change pas", 1)
add_bullet(doc, "Les annuaires et pages éditoriales restent publics.")
add_bullet(doc, "Premium ne donne pas automatiquement les droits de gestion d’un organisateur.")
add_bullet(doc, "La publication directe d’événements reste liée au statut Organisateur.")

# Page 11
doc.add_page_break()
add_title(doc, "Compte Organisateur", "Gestion professionnelle")
add_intro(doc, "L’offre Organisateur est destinée aux associations, organisateurs de soirées, clubs, artistes ou structures qui veulent gérer leur présence et publier directement.")
add_heading(doc, "Fonctions principales", 1)
for t in [
    "Réclamer une fiche existante ou créer une nouvelle fiche professionnelle.",
    "Modifier la présentation, la ville, les liens, les contacts, le logo et l’identité visuelle.",
    "Publier directement des événements et les rattacher à la bonne structure.",
    "Publier des actualités sur la fiche, avec photo, puis les modifier ou les supprimer.",
    "Suivre les abonnés, événements, intérêts, commentaires, rappels et actualités.",
    "Utiliser la messagerie et publier dans la Communauté au nom de l’organisme.",
    "Gérer une équipe lorsque plusieurs personnes doivent intervenir sur la fiche.",
]:
    add_bullet(doc, t)
add_heading(doc, "Offres professionnelles affichées", 1)
table = doc.add_table(rows=4, cols=3)
set_table_width(table, [2550, 1700, 5110])
table.style = "Table Grid"
for i, h in enumerate(("Offre", "Tarif", "Objet")):
    shade(table.cell(0, i), (GREEN, RED, GOLD)[i])
    set_cell_text(table.cell(0, i), h, bold=True, color=WHITE, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
for i, row in enumerate([
    ("Offre Organisateur", "15 € / mois", "Fiche, publication directe, profil, abonnés, messages et visibilité renforcée."),
    ("Orga Pro", "15 € / mois\nou 149 € / an", "Rappels J-3, calendrier intégrable, événements récurrents, statistiques avancées et un boost mensuel annoncés."),
    ("Boost événement", "24 € / 7 jours", "Événement à la une, bannière d’accueil et rappel aux personnes intéressées."),
], start=1):
    for j, text in enumerate(row):
        shade(table.cell(i, j), WHITE if i % 2 else LIGHT)
        set_cell_text(table.cell(i, j), text, bold=(j == 0), size=8.7)
add_callout(doc, "État actuel", "Les statistiques de base existent. Les statistiques avancées, le calendrier intégrable, les récurrences automatiques, les boosts mensuels et le paiement sécurisé sont annoncés mais peuvent nécessiter une activation ultérieure.", PALE_GOLD, GOLD)

# Page 12
doc.add_page_break()
add_title(doc, "État des fonctionnalités", "Disponible ou à finaliser")
add_intro(doc, "Cette dernière page sert de référence rapide pour la communication avec les membres, partenaires et organisateurs.")
table = doc.add_table(rows=1, cols=3)
set_table_width(table, [2650, 1910, 4800])
table.style = "Table Grid"
for i, h in enumerate(("Fonction", "Statut", "Précision")):
    shade(table.cell(0, i), INK)
    set_cell_text(table.cell(0, i), h, bold=True, color=WHITE, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
rows = [
    ("Agenda, recherche, filtres, calendrier", "DISPONIBLE", "Public et indexable."),
    ("Annuaires, cartes et fiches", "DISPONIBLE", "Publics ; contacts selon les données disponibles."),
    ("Compte, profil, pseudo unique", "DISPONIBLE", "Connexion requise pour les données personnelles."),
    ("Communauté, commentaires, messages", "DISPONIBLE", "Réservé aux membres connectés."),
    ("Blocage utilisateur", "DISPONIBLE", "Protection contre les échanges indésirables."),
    ("Notifications et suivi", "DISPONIBLE", "Membres et organisateurs."),
    ("Publication directe organisateur", "DISPONIBLE", "Après rattachement de la fiche et activation du rôle."),
    ("Statistiques organisateur de base", "DISPONIBLE", "Événements, intérêts, commentaires, rappels, abonnés, actus."),
    ("Paiement en ligne", "BIENTÔT", "Liens de paiement encore non configurés."),
    ("Billetterie interne", "BIENTÔT", "Les liens externes restent utilisables."),
    ("Avantages partenaires Premium", "VARIABLE", "Disponibles seulement lorsqu’un partenaire est actif."),
    ("Statistiques avancées / boosts Pro", "BIENTÔT", "Fonctions annoncées dans l’offre professionnelle."),
]
for idx, row in enumerate(rows):
    cells = table.add_row().cells
    status_color = GREEN if row[1] == "DISPONIBLE" else GOLD if row[1] == "VARIABLE" else RED
    for j, text in enumerate(row):
        shade(cells[j], WHITE if idx % 2 == 0 else LIGHT)
        set_cell_text(cells[j], text, bold=(j in (0, 1)), color=status_color if j == 1 else INK, size=8.4)
add_callout(doc, "Résumé", "Le site est déjà exploitable pour la découverte publique, la vie communautaire et la gestion de base des organisateurs. Les prochaines étapes concernent surtout la monétisation, les partenariats Premium et les outils Pro avancés.", PALE_GREEN, GREEN)

doc.core_properties.title = "Guide des fonctionnalités et des niveaux d’accès — Malagasy Events"
doc.core_properties.subject = "Public, Inscrit, Premium et Organisateur"
doc.core_properties.author = "Malagasy Events"
doc.core_properties.keywords = "Malagasy Events, fonctionnalités, accès, Premium, organisateur"
doc.save(OUT)
print(OUT)
